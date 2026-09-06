#!/usr/bin/env python3
"""
Generate professional client demo deck (.pptx) and presenter script (.docx)
for Aarogya One Connect client sales presentations.
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches as DocxInches, Pt as DocxPt, RGBColor as DocxRGBColor

ROOT = Path(__file__).resolve().parent

# Color Palette
DARK_TEAL = RGBColor(0x1E, 0x52, 0x4E)   # #1E524E
TEAL = RGBColor(0x2A, 0x6F, 0x6A)        # #2A6F6A
LIGHT_BG = RGBColor(0xF8, 0xFA, 0xFC)    # #F8FAFC
CARD_BG = RGBColor(0xFF, 0xFF, 0xFF)     # #FFFFFF
DARK_INK = RGBColor(0x0F, 0x17, 0x2A)    # #0F172A
MUTED_TEXT = RGBColor(0x47, 0x55, 0x69)  # #475569
ACCENT_GREEN = RGBColor(0x0D, 0x94, 0x88)# #0D9488


def create_pptx_deck() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 Widescreen
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Helper function to add background color
    def add_bg(slide, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.color.rgb = color
        return shape

    # Helper function to add header banner
    def add_header(slide, title_text, category="AAROGYA ONE CONNECT · DEMO DECK"):
        # Header background
        header_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.1))
        header_bg.fill.solid()
        header_bg.fill.fore_color.rgb = DARK_TEAL
        header_bg.line.color.rgb = DARK_TEAL

        # Header Text
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.15), Inches(11.733), Inches(0.8))
        tf = txBox.text_frame
        tf.word_wrap = True
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(9)
        p_cat.font.bold = True
        p_cat.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Helper function to add card container
    def add_card(slide, left, top, width, height, title, items, bg_color=CARD_BG, border_color=TEAL):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        txBox = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = txBox.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK_TEAL
        p_t.space_after = Pt(10)

        for item in items:
            p = tf.add_paragraph()
            p.text = "•  " + item
            p.font.size = Pt(12)
            p.font.color.rgb = DARK_INK
            p.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Dark Theme)
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1, DARK_TEAL)

    tx1 = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(3.5))
    tf1 = tx1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "Aarogya One Connect"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.space_after = Pt(12)

    p2 = tf1.add_paragraph()
    p2.text = "The Professional Outpatient (OPD) Operating System"
    p2.font.size = Pt(24)
    p2.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    p2.space_after = Pt(28)

    p3 = tf1.add_paragraph()
    p3.text = "Voice-to-Rx  ·  4-Role Desks  ·  UPI QR Billing  ·  Gynae & GP Workflows"
    p3.font.size = Pt(14)
    p3.font.bold = True
    p3.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)

    # Footer note
    tx1_ft = s1.shapes.add_textbox(Inches(1.0), Inches(6.2), Inches(11.333), Inches(0.8))
    p_ft = tx1_ft.text_frame.paragraphs[0]
    p_ft.text = "www.aarogyaoneconnect.in  |  sales@aarogyaoneconnect.in  |  +91 75660 99983"
    p_ft.font.size = Pt(12)
    p_ft.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)

    # -------------------------------------------------------------
    # SLIDE 2: The OPD Reality Today
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2, LIGHT_BG)
    add_header(s2, "The Modern OPD Challenge — Where Time & Revenue Slip Away")

    add_card(s2, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2), "Documentation Friction", [
        "Doctors spend 15–20 minutes typing or writing prescriptions manually every shift.",
        "Illegible handwritten Rx leads to patient confusion and follow-up errors.",
        "Manual repetition of patient history and vitals slows consultation speed."
    ])

    add_card(s2, Inches(6.8), Inches(1.5), Inches(5.6), Inches(5.2), "Fragmented Operations", [
        "Front desk, nurse, doctor, and lab operate in isolated silos.",
        "Chasing payment status at reception creates long queue waiting times.",
        "Risk of data leakage and non-compliance with DPDPA 2023 guidelines."
    ])

    # -------------------------------------------------------------
    # SLIDE 3: One OPD, Four Tailored Desks
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3, LIGHT_BG)
    add_header(s3, "Four Workspaces — Tailored for Every Team Member")

    col_w = Inches(2.7)
    gap = Inches(0.25)
    top_pos = Inches(1.5)
    h_pos = Inches(5.2)

    add_card(s3, Inches(0.8), top_pos, col_w, h_pos, "Doctor Desk", [
        "Voice-to-Rx dictation",
        "Case history brief",
        "Vitals trend graphs",
        "Video consultation",
        "Referral inbox"
    ])

    add_card(s3, Inches(0.8) + col_w + gap, top_pos, col_w, h_pos, "Staff / Nurse Desk", [
        "Patient check-in vitals",
        "Clinical history entry",
        "Symptom recording",
        "Pediatric ranges",
        "Offline queue support"
    ])

    add_card(s3, Inches(0.8) + (col_w + gap) * 2, top_pos, col_w, h_pos, "Reception Desk", [
        "Patient locking & MRN",
        "Live waiting queue",
        "Razorpay UPI QR",
        "Appointment desk",
        "Billing ledger"
    ])

    add_card(s3, Inches(0.8) + (col_w + gap) * 3, top_pos, col_w, h_pos, "Lab Desk", [
        "Diagnostic orders view",
        "Report document upload",
        "Structured result entry",
        "HL7 v2 ORU ingest",
        "Today's lab list"
    ])

    # -------------------------------------------------------------
    # SLIDE 4: Voice-to-Rx Mastery
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4, LIGHT_BG)
    add_header(s4, "Voice-to-Rx Engine — Dictate naturally, sign in seconds")

    add_card(s4, Inches(0.8), Inches(1.5), Inches(3.6), Inches(5.2), "1. Speak Naturally", [
        "Dictate consultation in English, Hindi, or auto-detect.",
        "Uses local/cloud AI speech engine.",
        "Audio processed in-memory (never stored on disk)."
    ])

    add_card(s4, Inches(4.8), Inches(1.5), Inches(3.6), Inches(5.2), "2. AI Scribe Parsing", [
        "LLM automatically structures symptoms, diagnosis, meds, and advice.",
        "Learns clinic-specific medical terminology and vocabulary.",
        "Highlights ANC & drug alert hints automatically."
    ])

    add_card(s4, Inches(8.8), Inches(1.5), Inches(3.6), Inches(5.2), "3. Signed PDF Rx", [
        "Instant review & 1-click digital signature.",
        "Generates clean, branded PDF prescription.",
        "Saves 15–20 minutes every single day!"
    ])

    # -------------------------------------------------------------
    # SLIDE 5: Specialty Clinical Depth (Gynecology & GP)
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5, LIGHT_BG)
    add_header(s5, "Built-In Clinical Depth — Gynecologist & Polyclinic Ready")

    add_card(s5, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2), "Obstetric & ANC Workflows", [
        "Full Obstetric Profile: LMP, EDD (Naegele's rule), GPLA, blood group, Rh.",
        "Gestational age calculation on patient chip.",
        "Pregnancy-aware vitals trend charts (GA x-axis).",
        "Automatic ANC rule alerts & ultrasound scan cadence hints."
    ])

    add_card(s5, Inches(6.8), Inches(1.5), Inches(5.6), Inches(5.2), "Polyclinic & GP Intelligence", [
        "Clinical Search: Search history by drug, diagnosis, or symptom across visits.",
        "Forward Case History: Handoff patient records to colleague doctors.",
        "Pediatric vitals range indicators.",
        "Bilingual Hindi/English patient outputs."
    ])

    # -------------------------------------------------------------
    # SLIDE 6: Integrated Billing & Revenue Integrity
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6, LIGHT_BG)
    add_header(s6, "Revenue Integrity — Razorpay Dynamic UPI QR")

    add_card(s6, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2), "Instant Dynamic UPI QR", [
        "Display dynamic Razorpay UPI QR code directly on patient visit card.",
        "Patient scans & pays via PhonePe, Google Pay, Paytm, or BHIM.",
        "Webhook automatically confirms payment in clinic ledger."
    ])

    add_card(s6, Inches(6.8), Inches(1.5), Inches(5.6), Inches(5.2), "100% Revenue Retained", [
        "Zero per-appointment booking commissions.",
        "Zero marketplace transaction deductions.",
        "Clear, itemized billing summary for consultation, dispensing, and lab."
    ])

    # -------------------------------------------------------------
    # SLIDE 7: Privacy-First Architecture
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7, LIGHT_BG)
    add_header(s7, "Privacy-First Architecture — DPDP 2023 Compliant")

    add_card(s7, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2), "HMAC Blind Patient Identity", [
        "Patient phone numbers and names are never stored in plain text on clinical records.",
        "HMAC-SHA256 hashing secures identity tokens.",
        "Encrypted rest-tier storage with PIN PBKDF2 hashing."
    ])

    add_card(s7, Inches(6.8), Inches(1.5), Inches(5.6), Inches(5.2), "Dedicated Indian Cloud Hosting", [
        "Data isolated per clinic tenant in Mumbai region.",
        "No ad SDKs or third-party tracking scripts.",
        "Role-based PIN access with full actor audit trail."
    ])

    # -------------------------------------------------------------
    # SLIDE 8: Value-Driven Subscription Plans
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8, LIGHT_BG)
    add_header(s8, "Simple, Value-Driven Subscriptions — No Surprise Fees")

    add_card(s8, Inches(0.8), Inches(1.5), Inches(3.6), Inches(5.2), "Starter (Solo)", [
        "₹1,999 / mo (₹19,990 / yr)",
        "1 Doctor + 3 Staff Seats",
        "Full Voice-to-Rx Engine",
        "Doctor, Staff, Reception, Lab desks",
        "UPI QR Billing & Android app"
    ])

    add_card(s8, Inches(4.8), Inches(1.5), Inches(3.6), Inches(5.2), "Clinic (Polyclinic)", [
        "₹5,499 / mo (₹54,990 / yr)",
        "Up to 5 Doctors + Unlimited Staff",
        "Obstetric & Gynae workflows",
        "Clinical Search & Analytics",
        "Forward Case Handoff inbox",
        "Priority Support Desk"
    ], bg_color=RGBColor(0xF0, 0xFD, 0xFA), border_color=ACCENT_GREEN)

    add_card(s8, Inches(8.8), Inches(1.5), Inches(3.6), Inches(5.2), "Network (Chain)", [
        "Custom (from ₹12,999 / mo)",
        "Multi-site management console",
        "SLA-backed uptime",
        "Custom branding on artifacts",
        "Historical data migration"
    ])

    # -------------------------------------------------------------
    # SLIDE 9: Go Live in 10 Minutes
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_bg(s9, LIGHT_BG)
    add_header(s9, "Go Live in 10 Minutes — Zero Complex Installation")

    add_card(s9, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2), "Mobile & Browser Flexibility", [
        "Android App (APK/Play Store) for doctors and mobile staff.",
        "Browser Web Desk (app.aarogyaoneconnect.in) for reception PC.",
        "Runs smoothly on existing clinic phones, tablets, and desktops."
    ])

    add_card(s9, Inches(6.8), Inches(1.5), Inches(5.6), Inches(5.2), "Seamless Onboarding", [
        "Self-onboarding takes under 10 minutes.",
        "14-Day Risk-Free Private Pilot available.",
        "Export your data anytime with zero lock-in."
    ])

    # -------------------------------------------------------------
    # SLIDE 10: Call to Action (Dark Theme)
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_bg(s10, DARK_TEAL)

    tx10 = s10.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(4.5))
    tf10 = tx10.text_frame
    tf10.word_wrap = True

    p = tf10.paragraphs[0]
    p.text = "Transform Your OPD Operations Today"
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.space_after = Pt(16)

    p2 = tf10.add_paragraph()
    p2.text = "Book your 14-day risk-free pilot or live product walkthrough."
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    p2.space_after = Pt(36)

    p3 = tf10.add_paragraph()
    p3.text = "📞 WhatsApp / Phone: +91 75660 99983"
    p3.font.size = Pt(18)
    p3.font.bold = True
    p3.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)
    p3.space_after = Pt(10)

    p4 = tf10.add_paragraph()
    p4.text = "✉️ Email: sales@aarogyaoneconnect.in"
    p4.font.size = Pt(18)
    p4.font.bold = True
    p4.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p4.space_after = Pt(10)

    p5 = tf10.add_paragraph()
    p5.text = "🌐 Web: www.aarogyaoneconnect.in"
    p5.font.size = Pt(18)
    p5.font.bold = True
    p5.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)

    pptx_path = ROOT / "Client_Demo_Deck.pptx"
    prs.save(str(pptx_path))
    return pptx_path


def create_presenter_script_docx() -> Path:
    doc = Document()
    
    # Page setup
    section = doc.sections[0]
    section.top_margin = DocxInches(0.8)
    section.bottom_margin = DocxInches(0.8)
    section.left_margin = DocxInches(0.8)
    section.right_margin = DocxInches(0.8)

    def set_font(run, size=11, bold=False, italic=False, color=None):
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        run.font.size = DocxPt(size)
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = color

    # Title
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Aarogya One Connect — Doctor Demo Script & Walkthrough")
    set_font(r, size=18, bold=True, color=DocxRGBColor(0x1E, 0x52, 0x4E))
    t.paragraph_format.space_after = DocxPt(2)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub.add_run("Slide-by-slide presenter notes for high-conversion clinic demos (Aug 2026)")
    set_font(r_sub, size=10, italic=True, color=DocxRGBColor(0x64, 0x74, 0x8B))
    sub.paragraph_format.space_after = DocxPt(16)

    slides_script = [
        ("Slide 1: Title & Welcome", 
         "Good morning Doctor / Team. Today I'm excited to show you Aarogya One Connect — an operating system built specifically for modern Indian outpatient clinics. Unlike legacy EMRs that feel like typing software, Aarogya bridges the gap between your phone and your front-desk PC.",
         "Show Slide 1; set a relaxed, professional tone.",
         "How many patients do you typically see in an OPD session?"),
        
        ("Slide 2: The OPD Reality Today", 
         "We know doctor time is precious. Right now, most practitioners spend 15 to 20 minutes per shift manually writing or typing prescriptions. Meanwhile, reception and nurses work in silos, and tracking UPI payments creates front-desk friction.",
         "Emphasize time loss and payment leakage.",
         "How much time does your team spend daily re-entering vitals or checking billing status?"),

        ("Slide 3: Four Workspaces, One Clinic", 
         "With Aarogya, everyone has a workspace tailored to their exact job. The doctor dictates notes on mobile; nurses record vitals; reception manages queues and UPI billing; the lab uploads diagnostic reports. Everyone stays synced in real time.",
         "Point to the 4 column cards.",
         "Would having dedicated views for reception and lab reduce noise in your consultation room?"),

        ("Slide 4: Voice-to-Rx Engine (Core Aha! Moment)", 
         "Here is our signature feature: Voice-to-Rx. You simply dictate your consultation notes naturally in Hindi, English, or mixed speech. Our AI engine automatically structures chief complaints, diagnoses, medications, and advice into a signed PDF in seconds.",
         "Perform live 15-second dictation demo on app if possible.",
         "Can you imagine dictating in Hindi or English and getting a printed signed Rx without typing?"),

        ("Slide 5: Specialty Depth (Gynae & GP)", 
         "If you practice Gynecology or Polyclinic consultations, Aarogya has deep clinical logic built-in: Obstetric cards with Naegele EDD calculations, pregnancy-aware vitals trends, ANC milestone alerts, and referral handoff between colleagues.",
         "Highlight Obstetric profile & ANC alerts.",
         "Do you see pregnant patients where tracking gestational age trends automatically would save mental energy?"),

        ("Slide 6: Integrated Billing & UPI QR", 
         "For billing, reception displays a dynamic Razorpay UPI QR directly on the patient's card. The patient scans via PhonePe or Google Pay, and the ledger updates instantly. Best of all: 100% revenue is yours — zero booking or transaction commissions.",
         "Show payment QR generation on patient card.",
         "How currently do you verify if a patient paid their consultation fee at reception?"),

        ("Slide 7: Privacy-First Architecture (DPDP)", 
         "We protect your practice and patients. Using HMAC Blind Identity, searchable phone numbers are never stored in plain text. Hosted on dedicated Indian cloud servers in Mumbai with strict role PIN access.",
         "Highlight privacy posture.",
         "Is patient data security and DPDP compliance something your clinic considers?"),

        ("Slide 8: Simple Subscription Plans", 
         "Our pricing is completely transparent. Starter plan is ₹1,999/mo with 1 doctor and 3 staff seats. Clinic plan for polyclinics is ₹5,499/mo for up to 5 doctors and unlimited staff. No hidden fees or surprise per-message bills.",
         "Show plan breakdown.",
         "Which structure fits your clinic better — solo or multi-doctor?"),

        ("Slide 9: Go Live in 10 Minutes", 
         "Setting up takes under 10 minutes. Android app for doctors, web browser desk for PCs. We offer a 14-day risk-free trial on real patients with full onboarding support.",
         "Mention fast onboarding.",
         "Shall we set up your 14-day trial tenant right now?"),

        ("Slide 10: Call to Action & Close", 
         "Thank you Doctor! Let's get your clinic started on Aarogya One Connect today.",
         "Hand over the 1-Page Walk-In Sales Sheet.",
         "What start date works best for your 14-day trial?")
    ]

    for title_str, script, action, question in slides_script:
        h = doc.add_paragraph()
        r_h = h.add_run(title_str)
        set_font(r_h, size=13, bold=True, color=DocxRGBColor(0x1E, 0x52, 0x4E))
        h.paragraph_format.space_before = DocxPt(10)
        h.paragraph_format.space_after = DocxPt(2)

        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_after = DocxPt(4)
        r_lbl = p_s.add_run("Verbal Script: ")
        set_font(r_lbl, size=10.5, bold=True, color=DocxRGBColor(0x2A, 0x6F, 0x6A))
        r_txt = p_s.add_run(f'"{script}"')
        set_font(r_txt, size=10.5, italic=True)

        p_a = doc.add_paragraph()
        p_a.paragraph_format.space_after = DocxPt(4)
        r_albl = p_a.add_run("Presenter Action: ")
        set_font(r_albl, size=10, bold=True, color=DocxRGBColor(0x47, 0x55, 0x69))
        r_atxt = p_a.add_run(action)
        set_font(r_atxt, size=10)

        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_after = DocxPt(8)
        r_qlbl = p_q.add_run("Discovery Question: ")
        set_font(r_qlbl, size=10, bold=True, color=DocxRGBColor(0x0D, 0x94, 0x88))
        r_qtxt = p_q.add_run(question)
        set_font(r_qtxt, size=10, bold=True)

    script_path = ROOT / "Presenter_Demo_Script.docx"
    doc.save(str(script_path))
    return script_path


if __name__ == "__main__":
    p_pptx = create_pptx_deck()
    p_docx = create_presenter_script_docx()
    print(f"Generated Presentation Assets:\n1. {p_pptx}\n2. {p_docx}")
